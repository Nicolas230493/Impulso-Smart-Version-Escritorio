"""Tests de la lógica de escritorio (business_logic + ticket_printer) sobre una base temporal.

Uso: python desktop_tests.py
Crea un db.sqlite3 vacío en una carpeta temporal, lo inicializa igual que el arranque real
(core.bootstrap.prepare_database: migraciones + datos base) y apunta SQLAlchemy a esa base:
nunca toca la base real. (No se llama test*.py para que `manage.py test`
no lo recoja: necesita SQLAlchemy y su propia base.)
"""
import ast
import datetime
import inspect
import os
import sys
import tempfile
import unittest
import uuid
import warnings
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'pos_system.settings')

TMP_DIR = tempfile.TemporaryDirectory()
TMP_DB = Path(TMP_DIR.name) / 'db.sqlite3'

import django
from django.conf import settings
settings.DATABASES['default']['NAME'] = TMP_DB  # antes de abrir cualquier conexión
django.setup()
from django.test.utils import setup_test_environment
setup_test_environment()  # permite usar el cliente de pruebas de Django (dashboard web) sobre esta base

from sqlalchemy.exc import SAWarning
warnings.simplefilter("error", SAWarning)  # cualquier aviso de SQLAlchemy (ej. Decimal en SQLite) es un error

import db_manager
import business_logic as bl
import report_export
import ticket_printer
from core import bootstrap
db_manager.db = bl.db = db_manager.DatabaseManager(TMP_DB)
BOOTSTRAP_CREATED = bootstrap.prepare_database(db_manager.db)  # base vacía -> esquema + datos base

from django.contrib.auth.models import User as DUser
from customers.models import Customer as DCustomer
from finance.models import PaymentMethod as DPaymentMethod
from products.models import Product as DProduct, StockLoss as DStockLoss
from sales.models import SaleReturn as DSaleReturn, SaleReturnDetail as DSaleReturnDetail
from suppliers.models import Supplier as DSupplier
from db_manager import (CashSession, CurrentAccount, Customer, InventoryMovement, Product, Purchase, PurchaseDetail,
                        Sale, SaleDetail, SalePayment)

# Medios sembrados por el bootstrap
PM = {m.name: m.id for m in DPaymentMethod.objects.all()}
# Como en el arranque real (login_gui.check_admin_exists): el usuario 1 es el admin que usa el escritorio
DUser.objects.get_or_create(id=1, defaults={"username": "admin"})


def make_product(name="Prod", weighable=False, stock="10", price="1000", cost="500", barcode=None):
    return DProduct.objects.create(sku=f"T-{uuid.uuid4().hex[:8]}", barcode=barcode, name=name, unit="kg" if weighable else "unidad",
                                   price=D(price), cost_price=D(cost), stock=D(stock), is_weighable=weighable).id


def make_customer(balance="0", limit="0"):
    return DCustomer.objects.create(full_name=f"Cliente {uuid.uuid4().hex[:6]}", dni_cuit=uuid.uuid4().hex[:11],
                                    balance=D(balance), limite_credito=D(limit)).id


def get(model, pk):
    s = db_manager.db.get_session()
    try:
        return s.get(model, pk)
    finally:
        s.close()


def open_cash():
    s = db_manager.db.get_session()
    try:
        return s.query(CashSession).filter_by(is_open=True, user_id=1).one()
    finally:
        s.close()


class SchemaTests(unittest.TestCase):
    def test_schema_matches_sqlalchemy(self):
        self.assertEqual(db_manager.db.missing_columns(), {})


class SchemaSyncTests(unittest.TestCase):
    """Regresión: SQLAlchemy debe espejar exactamente las tablas de Django (ej. ActivityLog.date vs timestamp)."""
    KIND = {'IntegerField': 'int', 'BigAutoField': 'int', 'AutoField': 'int', 'ForeignKey': 'int', 'PositiveIntegerField': 'int',
            'DecimalField': 'dec', 'CharField': 'str', 'TextField': 'str', 'EmailField': 'str', 'GenericIPAddressField': 'str',
            'BooleanField': 'bool', 'DateTimeField': 'dt', 'DateField': 'date'}

    @staticmethod
    def sa_kind(t):
        from sqlalchemy import Boolean, Date, DateTime, Integer, String, Text
        for cls, kind in ((Boolean, 'bool'), (db_manager.SqliteDecimal, 'dec'), (Integer, 'int'), (String, 'str'),
                          (Text, 'str'), (DateTime, 'dt'), (Date, 'date')):
            if isinstance(t, cls):
                return kind
        return type(t).__name__

    def test_mappings_match_django_models(self):
        import sqlite3
        from django.apps import apps
        models = {m._meta.db_table: m for m in apps.get_models()}
        con = sqlite3.connect(TMP_DB)
        problems = []
        for mapper in db_manager.Base.registry.mappers:
            table = mapper.local_table
            model = models.get(table.name)
            if not model:
                problems.append(f"{table.name}: no es una tabla de Django")
                continue
            fields = {f.column: f for f in model._meta.concrete_fields}
            for col in table.columns:
                field = fields.get(col.name)
                if not field:
                    problems.append(f"{table.name}.{col.name}: no existe en Django")
                    continue
                if self.KIND.get(field.get_internal_type()) != self.sa_kind(col.type):
                    problems.append(f"{table.name}.{col.name}: {field.get_internal_type()} vs {type(col.type).__name__}")
                if field.get_internal_type() == 'DecimalField' and col.type.impl.scale != field.decimal_places:
                    problems.append(f"{table.name}.{col.name}: decimales {field.decimal_places} vs {col.type.impl.scale}")
            # Columnas NOT NULL sin default en la base deben estar mapeadas (si no, el INSERT falla)
            mapped = {c.name for c in table.columns}
            for _, name, _, notnull, default, pk in con.execute(f'PRAGMA table_info("{table.name}")'):
                if notnull and default is None and not pk and name not in mapped:
                    problems.append(f"{table.name}.{name}: NOT NULL sin mapear")
        con.close()
        self.assertEqual(problems, [])


class GuiCallSignatureTests(unittest.TestCase):
    """Sin abrir Tk: cada llamada de gui_main a la capa de negocio respeta la firma real."""

    def test_gui_calls_match_signatures(self):
        modules = {"business_logic": bl, "ticket_printer": ticket_printer, "report_export": report_export}
        tree = ast.parse(Path(__file__).with_name("gui_main.py").read_text(encoding="utf-8"))
        imported = {}
        for node in tree.body:
            if isinstance(node, ast.ImportFrom) and node.module in modules:
                for alias in node.names:
                    imported[alias.asname or alias.name] = getattr(modules[node.module], alias.name)  # falla si no existe
        problems, checked = [], 0
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and callable(imported.get(node.func.id)):
                fn = imported[node.func.id]
                if inspect.isclass(fn) or any(isinstance(a, ast.Starred) for a in node.args):
                    continue
                if any(k.arg is None for k in node.keywords):
                    continue
                try:
                    inspect.signature(fn).bind(*[None] * len(node.args), **{k.arg: None for k in node.keywords})
                    checked += 1
                except TypeError as e:
                    problems.append(f"línea {node.lineno}: {node.func.id}(): {e}")
        self.assertEqual(problems, [])
        self.assertGreater(checked, 25)


class BootstrapTests(unittest.TestCase):
    def test_empty_database_is_seeded(self):
        # La migración core.0003 ya siembra; el paso posterior del bootstrap no tiene nada que crear
        expected = {name for name, _, _ in bootstrap.BASE_PAYMENT_METHODS}
        self.assertEqual(BOOTSTRAP_CREATED, [])
        self.assertEqual({m["name"] for m in bl.get_payment_methods()}, expected)
        cf = [c for c in bl.get_customers_list() if c["is_final_consumer"]]
        self.assertEqual(len(cf), 1)
        self.assertEqual(cf[0]["limite_credito"], D("0.00"))

    def test_second_run_is_idempotent(self):
        before = len(bl.get_payment_methods()), len(bl.get_customers_list())
        self.assertEqual(bootstrap.prepare_database(db_manager.db), [])
        self.assertEqual((len(bl.get_payment_methods()), len(bl.get_customers_list())), before)


class NetProfitTests(unittest.TestCase):
    """Ganancia Neta Real del turno = Ventas netas - CMV - Gastos - Mermas valorizadas."""

    @classmethod
    def setUpClass(cls):
        # Cajero propio para aislar el turno de las ventas de otros tests
        cls.user_id = DUser.objects.create_user(f"cajero{uuid.uuid4().hex[:6]}").id
        ok, msg = bl.open_cash_session(D("1000"), user_id=cls.user_id)
        assert ok, msg
        pollo = make_product("Pollo", weighable=True, stock="1", price="4000", cost="2000")
        coca = make_product("Coca", stock="10", price="1500", cost="900")
        ok, res = bl.process_sale([{"id": pollo, "price": 4000, "quantity": "1,250"}, {"id": coca, "price": 1500, "quantity": 2}],
                                  None, PM["Efectivo"], user_id=cls.user_id)  # $8000, CMV 2500 + 1800
        assert ok, res
        cls.sale_id = res["sale_id"]
        ok, msg = bl.add_cash_expense(D("500"), "Hielo", user_id=cls.user_id)
        assert ok, msg
        # Devolución y mermas se cargan desde la web (Django): mismas fechas UTC que el escritorio
        ret = DSaleReturn.objects.create(sale_id=cls.sale_id, total_amount=D("1500"), user_id=cls.user_id)
        DSaleReturnDetail.objects.create(sale_return=ret, product_id=coca, quantity=D("1"), price_at_return=D("1500"))
        DStockLoss.objects.create(product_id=pollo, quantity=D("0.350"), reason="DAM", unit_cost=D("2000"))
        DStockLoss.objects.create(product_id=coca, quantity=D("1"), reason="EXP")  # sin costo: usa el vigente (900)
        cls.session_id = bl.get_open_cash_session(user_id=cls.user_id)["id"]
        ok, msg = bl.close_cash_session(D("8000"), user_id=cls.user_id)
        assert ok, msg

    def test_net_profit_for_closed_session(self):
        s = bl.get_profitability_and_audit_summary(cash_session_id=self.session_id)
        self.assertEqual((s["gross_sales"], s["returns"], s["total_sales"]), (D("8000.00"), D("1500.00"), D("6500.00")))
        self.assertEqual(s["cmv"], D("3400.00"))
        self.assertEqual(s["gross_profit"], D("3100.00"))
        self.assertEqual(s["expenses"], D("500.00"))
        self.assertEqual(s["shrinkage_cost"], D("1600.00"))
        self.assertEqual(s["net_profit"], D("1000.00"))
        self.assertEqual(s["auto_adjustments_count"], 1)
        self.assertEqual(s["auto_adjustments"][0]["quantity"], D("0.250"))
        self.assertEqual(s["payments_breakdown"], {"Efectivo": D("8000.00")})
        self.assertEqual(bl.get_last_closed_cash_session(user_id=self.user_id)["id"], self.session_id)

    def test_range_summary_uses_local_dates(self):
        import datetime
        now = datetime.datetime.now()
        s = bl.get_profitability_and_audit_summary(now - datetime.timedelta(hours=1), now + datetime.timedelta(minutes=1))
        self.assertGreaterEqual(s["sales_count"], 1)
        self.assertGreaterEqual(s["shrinkage_cost"], D("1600.00"))
        self.assertLess(abs(s["start"] - (now - datetime.timedelta(hours=1))), datetime.timedelta(seconds=1))
        with self.assertRaises(ValueError):
            bl.get_profitability_and_audit_summary()

    def test_export_csv_and_excel(self):
        from openpyxl import load_workbook
        s = bl.get_profitability_and_audit_summary(cash_session_id=self.session_id)
        xlsx = report_export.export_summary(s, Path(TMP_DIR.name) / "resumen.xlsx")
        wb = load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["Resumen", "Pagos", "Despostes", "Mermas"])
        rows = {r[0]: r[1] for r in wb["Resumen"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["GANANCIA NETA REAL"], 1000.0)
        self.assertEqual(wb["Despostes"].max_row, 2)
        self.assertEqual(wb["Mermas"].max_row, 3)
        csv_path = report_export.export_summary(s, Path(TMP_DIR.name) / "resumen.csv")
        text = csv_path.read_text(encoding="utf-8-sig")
        self.assertIn("GANANCIA NETA REAL;1000.00", text)
        self.assertIn("DESPOSTES", text)


class SaleTests(unittest.TestCase):
    def test_weighable_sale_with_auto_adjust(self):
        asado = make_product("Asado", weighable=True, stock="0.300", price="8999.99")
        coca = make_product("Coca", stock="10", price="1500")
        ok, res = bl.process_sale([{"id": asado, "price": D("8999.99"), "quantity": "1,250"},
                                   {"id": coca, "price": 1500, "quantity": 2}], None, PM["Efectivo"])
        self.assertTrue(ok, res)
        self.assertEqual(res["total"], D("14249.99"))
        self.assertEqual([a["adjusted_qty"] for a in res["auto_adjustments"]], [D("0.950")])
        self.assertEqual(bl.to_decimal_qty(get(Product, asado).stock), D("0.000"))
        self.assertEqual(bl.to_decimal_qty(get(Product, coca).stock), D("8.000"))
        s = db_manager.db.get_session()
        movs = sorted((m.movement_type, bl.to_decimal_qty(m.quantity)) for m in s.query(InventoryMovement).filter_by(product_id=asado))
        s.close()
        self.assertEqual(movs, [("ADJ", D("0.950")), ("OUT", D("1.250"))])

    def test_invalid_line_rolls_back_whole_sale(self):
        pollo = make_product("Pollo", weighable=True, stock="5")
        coca = make_product("Coca")
        ok, err = bl.process_sale([{"id": pollo, "price": 100, "quantity": "0.5"},
                                   {"id": coca, "price": 100, "quantity": "1.5"}], None, PM["Efectivo"])
        self.assertFalse(ok)
        self.assertIn("entera", err)
        self.assertEqual(bl.to_decimal_qty(get(Product, pollo).stock), D("5.000"))


class MixedPaymentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        bl.open_cash_session(D("500"))

    def setUp(self):
        self.pollo = make_product("Pollo", weighable=True, stock="20", price="4000")
        self.cart = [{"id": self.pollo, "price": D("4000"), "quantity": "1,250"}]  # $5000.00

    def test_payments_must_match_total(self):
        ok, err = bl.process_sale(self.cart, None, payments=[{"payment_method_id": PM["Efectivo"], "amount": "2000"}])
        self.assertFalse(ok)
        self.assertIn("no coincide", err)

    def test_credit_limit_checked_on_credit_portion_only(self):
        cust = make_customer(balance="1000", limit="5000")
        ok, err = bl.process_sale(self.cart, cust, payments=[{"payment_method_id": PM["Efectivo"], "amount": 500},
                                                             {"payment_method_id": PM["Cuenta Corriente"], "amount": 4500}])
        self.assertFalse(ok)
        self.assertIn("Límite", err)

    def test_final_consumer_cannot_take_credit(self):
        cf = DCustomer.objects.get(dni_cuit=bl.FINAL_CONSUMER_DNI).id
        ok, err = bl.process_sale(self.cart, cf, payments=[{"payment_method_id": PM["Efectivo"], "amount": 4000},
                                                           {"payment_method_id": PM["Cuenta Corriente"], "amount": 1000}])
        self.assertFalse(ok)

    def test_mixed_sale_splits_cash_digital_and_credit(self):
        cust = make_customer(balance="1000", limit="5000")
        before = open_cash()
        ok, res = bl.process_sale(self.cart, cust, payments=[
            {"payment_method_id": PM["Efectivo"], "amount": "2000"},
            {"payment_method_id": PM["Transferencia / QR"], "amount": "1500,00"},
            {"payment_method_id": PM["Cuenta Corriente"], "amount": D("1500")}])
        self.assertTrue(ok, res)
        after = open_cash()
        self.assertEqual(bl.to_decimal_money(after.total_sales_cash - before.total_sales_cash), D("2000.00"))
        self.assertEqual(bl.to_decimal_money(after.total_sales_digital - before.total_sales_digital), D("1500.00"))
        self.assertEqual(bl.to_decimal_money(get(Customer, cust).balance), D("2500.00"))
        self.assertEqual(get(Sale, res["sale_id"]).payment_method_id, PM["Efectivo"])
        s = db_manager.db.get_session()
        pays = sorted((p.payment_method_id, bl.to_decimal_money(p.amount)) for p in s.query(SalePayment).filter_by(sale_id=res["sale_id"]))
        debt = s.query(CurrentAccount).filter_by(customer_id=cust).one()
        s.close()
        self.assertEqual(pays, sorted([(PM["Efectivo"], D("2000.00")), (PM["Transferencia / QR"], D("1500.00")), (PM["Cuenta Corriente"], D("1500.00"))]))
        self.assertEqual(bl.to_decimal_money(debt.amount), D("1500.00"))

    def test_single_payment_method_is_backwards_compatible(self):
        ok, res = bl.process_sale(self.cart, None, PM["Tarjeta Débito"])
        self.assertTrue(ok, res)
        self.assertEqual([(p["name"], p["amount"]) for p in res["payments"]], [("Tarjeta Débito", D("5000.00"))])


class PurchaseCPPTests(unittest.TestCase):
    def test_purchase_applies_cpp_and_registers_purchase(self):
        pollo = make_product("Pollo", weighable=True, stock="10", cost="2000")
        vacio = make_product("Vacío", weighable=True, stock="-1.2", cost="7777")
        supplier = DSupplier.objects.create(name="Avícola").id
        ok, res = bl.register_purchase_or_stock_entry([
            {"product_id": pollo, "quantity": "18,500", "unit_cost": "2600"},
            {"product_id": vacio, "quantity": "5", "unit_cost": "8000"}], supplier_id=supplier, invoice_number="A-1")
        self.assertTrue(ok, res)
        p, v = get(Product, pollo), get(Product, vacio)
        self.assertEqual((bl.to_decimal_qty(p.stock), bl.to_decimal_money(p.cost_price)), (D("28.500"), D("2389.47")))
        self.assertEqual((bl.to_decimal_qty(v.stock), bl.to_decimal_money(v.cost_price)), (D("3.800"), D("8000.00")))
        self.assertEqual(bl.to_decimal_money(get(Purchase, res["purchase_id"]).total_amount), D("88100.00"))
        s = db_manager.db.get_session()
        self.assertEqual(s.query(PurchaseDetail).filter_by(purchase_id=res["purchase_id"]).count(), 2)
        s.close()

    def test_failed_item_rolls_back_entry(self):
        pollo = make_product("Pollo", weighable=True, stock="10", cost="2000")
        ok, err = bl.register_purchase_or_stock_entry([{"product_id": pollo, "quantity": "1", "unit_cost": "1"},
                                                       {"product_id": 999999, "quantity": 1, "unit_cost": 1}])
        self.assertFalse(ok)
        p = get(Product, pollo)
        self.assertEqual((bl.to_decimal_qty(p.stock), bl.to_decimal_money(p.cost_price)), (D("10.000"), D("2000.00")))


class LookupTests(unittest.TestCase):
    def test_find_product_by_barcode_or_sku(self):
        pid = make_product("Yerba", barcode="7790001112223")
        self.assertEqual(bl.find_product_by_code(" 7790001112223 ")["id"], pid)
        self.assertEqual(bl.find_product_by_code(DProduct.objects.get(pk=pid).sku)["id"], pid)
        self.assertIsNone(bl.find_product_by_code("no-existe"))
        self.assertIsNone(bl.find_product_by_code(""))


class TicketTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        asado = make_product("Asado de tira premium envasado al vacío", weighable=True, stock="5", price="8999.99")
        coca = make_product("Coca", stock="10", price="1500")
        cust = make_customer(limit="0")
        DCustomer.objects.filter(pk=cust).update(full_name="Juan Test")
        ok, res = bl.process_sale([{"id": asado, "price": D("8999.99"), "quantity": "1,250"},
                                   {"id": coca, "price": 1500, "quantity": 2}], cust, payments=[
            {"payment_method_id": PM["Efectivo"], "amount": "10000"},
            {"payment_method_id": PM["Transferencia / QR"], "amount": "4249.99"}])
        assert ok, res
        cls.sale_id = res["sale_id"]

    def test_ticket_content_80mm(self):
        text = ticket_printer.generate_sale_ticket_text(self.sale_id, 42)
        self.assertIn("IMPULSO SMART", text)
        self.assertIn(f"Venta N°: {self.sale_id:06d}", text)
        self.assertIn("Cliente: Juan Test", text)
        self.assertIn("1.250 kg x $8999.99", text)
        self.assertIn("2 u. x $1500.00", text)
        lines = text.splitlines()
        self.assertTrue(any(l.startswith("TOTAL") and l.endswith("$14249.99") for l in lines))
        self.assertTrue(any(l.strip().startswith("Efectivo") and l.endswith("$10000.00") for l in lines))
        self.assertTrue(any(l.strip().startswith("Transferencia") and l.endswith("$4249.99") for l in lines))

    def test_ticket_respects_width(self):
        for width in ticket_printer.TICKET_WIDTHS.values():
            text = ticket_printer.generate_sale_ticket_text(self.sale_id, width)
            self.assertTrue(all(len(l) <= width for l in text.splitlines()), width)
            self.assertIn("$11249.99", text)

    def test_save_ticket(self):
        text = ticket_printer.generate_sale_ticket_text(self.sale_id)
        path = ticket_printer.save_ticket(self.sale_id, text, folder=Path(TMP_DIR.name) / "tickets")
        self.assertEqual(path.name, f"ticket_venta_{self.sale_id:06d}.txt")
        self.assertEqual(path.read_text(encoding="utf-8"), text)

    def test_unknown_sale(self):
        with self.assertRaises(ValueError):
            ticket_printer.generate_sale_ticket_text(999999)


class GuiDataTests(unittest.TestCase):
    def test_gui_queries_return_plain_data(self):
        make_product("Dato")
        for rows in (bl.get_inventory_list(), bl.search_products_for_sale("Dato"), bl.get_customers_list(),
                     bl.get_suppliers_list(), bl.get_payment_methods()):
            self.assertIsInstance(rows, list)
            self.assertTrue(all(isinstance(r, dict) for r in rows))
        top = bl.get_bi_stats()["top_products"]
        self.assertTrue(all({"name", "quantity", "is_weighable"} <= set(t) for t in top))



class ProductEditingTests(unittest.TestCase):
    """Regresión: el escritorio no podía crear productos (created_at/updated_at sin mapear)."""

    def test_create_product_with_initial_stock_and_expiry(self):
        ok, pid = bl.save_product({"name": "Queso", "price": "5500,50", "cost_price": "3000", "is_weighable": True,
                                   "unit": "unidad", "min_stock": "1,5", "tax_rate": "10.5", "expiry_date": "31/12/2026",
                                   "barcode": "7791234567890", "initial_stock": "2,750"})
        self.assertTrue(ok, pid)
        p = DProduct.objects.get(pk=pid)  # Django lee lo que escribió el escritorio
        self.assertEqual((p.unit, p.stock, p.price, p.min_stock), ("kg", D("2.750"), D("5500.50"), D("1.500")))
        self.assertEqual(p.expiry_date, datetime.date(2026, 12, 31))
        self.assertTrue(p.sku.startswith("SKU-"))
        self.assertIsNotNone(p.created_at)
        self.assertTrue(p.movements.filter(movement_type="IN", reference="Stock inicial").exists())

    def test_edit_logs_price_change_and_validates(self):
        ok, pid = bl.save_product({"name": "Yerba", "price": "100", "sku": "YER-1"})
        self.assertTrue(ok, pid)
        ok, res = bl.save_product({"name": "Yerba 1kg", "price": "120", "sku": "YER-1"}, product_id=pid)
        self.assertTrue(ok, res)
        self.assertEqual(DProduct.objects.get(pk=pid).price_logs.get().new_price, D("120.00"))
        self.assertEqual(bl.save_product({"name": "Otro", "price": "1", "sku": "YER-1"}), (False, "Ya existe un producto con SKU 'YER-1'."))
        self.assertFalse(bl.save_product({"name": "", "price": "1"})[0])
        self.assertFalse(bl.save_product({"name": "X", "price": "-1"})[0])
        self.assertFalse(bl.save_product({"name": "X", "price": "1", "expiry_date": "31-31-2026"})[0])
        self.assertFalse(bl.save_product({"name": "X", "price": "1", "initial_stock": "1,5"})[0])  # unitario fraccional


class HardeningTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.user_id = DUser.objects.create_user(f"hard{uuid.uuid4().hex[:6]}").id

    def test_empty_cart_and_unknown_customer(self):
        self.assertEqual(bl.process_sale([], None, PM["Efectivo"]), (False, "El carrito está vacío."))
        pid = make_product("Pan")
        ok, err = bl.process_sale([{"id": pid, "price": 10, "quantity": 1}], 999999, PM["Efectivo"])
        self.assertFalse(ok)
        self.assertIn("inexistente", err)

    def test_zero_total_sale_needs_no_payment(self):
        pid = make_product("Muestra gratis", price="0")
        ok, res = bl.process_sale([{"id": pid, "price": 0, "quantity": 1}], None, PM["Efectivo"])
        self.assertTrue(ok, res)
        self.assertEqual((res["total"], res["payments"]), (D("0.00"), []))

    def test_cpp_with_missing_cost(self):
        self.assertEqual(bl.calculate_weighted_average_cost("10", None, "10", "100"), D("50.00"))
        self.assertEqual(bl.calculate_weighted_average_cost(None, None, "5", "100"), D("100.00"))

    def test_cash_session_validations(self):
        uid = self.user_id
        self.assertFalse(bl.open_cash_session("-1", user_id=uid)[0])
        self.assertFalse(bl.open_cash_session("abc", user_id=uid)[0])
        self.assertTrue(bl.open_cash_session("1000,50", user_id=uid)[0])
        self.assertFalse(bl.add_cash_expense("0", "Hielo", user_id=uid)[0])
        self.assertFalse(bl.add_cash_expense("100", "   ", user_id=uid)[0])
        self.assertTrue(bl.add_cash_expense("100,25", "Hielo", user_id=uid)[0])
        self.assertFalse(bl.close_cash_session("-5", user_id=uid)[0])
        ok, msg = bl.close_cash_session("900,25", user_id=uid)  # sin ventas: esperado 1000.50 - 100.25
        self.assertTrue(ok, msg)
        self.assertEqual(msg, "Caja cerrada. Diferencia: $0.00")
        summary = bl.get_profitability_and_audit_summary(cash_session_id=bl.get_last_closed_cash_session(user_id=uid)["id"])
        self.assertEqual((summary["sales_count"], summary["gross_margin_pct"], summary["net_profit"]), (0, D("0.0"), D("-100.25")))


class UserIsolationTests(unittest.TestCase):
    """Regresión: cada cajero opera sobre su propia caja (antes todo iba a user_id=1)."""

    def test_sale_impacts_only_the_sellers_cash_session(self):
        a = DUser.objects.create_user(f"a{uuid.uuid4().hex[:6]}").id
        b = DUser.objects.create_user(f"b{uuid.uuid4().hex[:6]}").id
        bl.open_cash_session("0", user_id=a)
        bl.open_cash_session("0", user_id=b)
        pid = make_product("Coca", price="1500")
        ok, res = bl.process_sale([{"id": pid, "price": 1500, "quantity": 1}], None, PM["Efectivo"], user_id=b)
        self.assertTrue(ok, res)
        self.assertEqual(get(Sale, res["sale_id"]).user_id, b)
        self.assertEqual(bl.get_open_cash_session(user_id=a)["total_sales_cash"], D("0.00"))
        self.assertEqual(bl.get_open_cash_session(user_id=b)["total_sales_cash"], D("1500.00"))
        self.assertFalse(bl.close_cash_session("0", user_id=999999)[0])  # sin caja propia no cierra la de otro


class TimezoneTests(unittest.TestCase):
    """Regresión: el día del negocio va de 00:00 a 23:59 hora argentina (no se pierden las 21–24 h)."""

    def test_local_day_includes_late_night_and_matches_django(self):
        from core.timeutils import BUSINESS_TIME_ZONE, local_to_utc
        from sales.models import Sale as DSale
        self.assertEqual(BUSINESS_TIME_ZONE, "America/Argentina/Buenos_Aires")
        pid = make_product("Fernet", price="100", cost="50")
        ids = []
        for _ in range(2):
            ok, res = bl.process_sale([{"id": pid, "price": 100, "quantity": 1}], None, PM["Efectivo"])
            self.assertTrue(ok, res)
            ids.append(res["sale_id"])
        # 15/01/2020 22:30 local (= 16/01 01:30 UTC) y 16/01/2020 00:30 local
        s = db_manager.db.get_session()
        s.get(Sale, ids[0]).date = local_to_utc(datetime.datetime(2020, 1, 15, 22, 30))
        s.get(Sale, ids[1]).date = local_to_utc(datetime.datetime(2020, 1, 16, 0, 30))
        s.commit()
        s.close()
        day = bl.get_profitability_and_audit_summary(datetime.datetime(2020, 1, 15), datetime.datetime(2020, 1, 15, 23, 59, 59, 999999))
        self.assertEqual((day["sales_count"], day["total_sales"]), (1, D("100.00")))
        self.assertEqual(day["start"], datetime.datetime(2020, 1, 15))
        # Django ve el mismo día del negocio
        self.assertEqual(list(DSale.objects.filter(date__date=datetime.date(2020, 1, 15)).values_list("id", flat=True)), [ids[0]])
        self.assertEqual(bl.get_sale_ticket_data(ids[0])["date"], datetime.datetime(2020, 1, 15, 22, 30))



class ShrinkageTests(unittest.TestCase):
    def test_register_shrinkage_values_at_current_cost(self):
        from products.models import LOSS_REASONS as DJANGO_REASONS
        self.assertEqual(bl.LOSS_REASONS, dict(DJANGO_REASONS))  # paridad con Django
        pid = make_product("Vacío", weighable=True, stock="5", cost="7000")
        ok, msg, data = bl.register_shrinkage(pid, "0,350", "Descarte / Merma", user_id=1, notes="desposte")
        self.assertTrue(ok, msg)
        self.assertEqual((data["quantity"], data["unit_cost"], data["cost"], data["reason"]), (D("0.350"), D("7000.00"), D("2450.00"), "MER"))
        self.assertEqual(bl.to_decimal_qty(get(Product, pid).stock), D("4.650"))
        loss = DStockLoss.objects.get(pk=data["loss_id"])  # Django lee la baja del escritorio
        self.assertEqual((loss.unit_cost, loss.get_reason_display(), loss.notes), (D("7000.00"), "Descarte / Merma", "desposte"))
        self.assertTrue(DProduct.objects.get(pk=pid).movements.filter(movement_type="OUT", reference="BAJA: MER").exists())

    def test_free_text_reason_and_validations(self):
        pid = make_product("Yogur", stock="3", cost="500")
        ok, msg, data = bl.register_shrinkage(pid, 1, "Se cortó la cadena de frío")
        self.assertTrue(ok, msg)
        self.assertEqual((data["reason"], DStockLoss.objects.get(pk=data["loss_id"]).notes), ("OTH", "Se cortó la cadena de frío"))
        self.assertFalse(bl.register_shrinkage(pid, "0,5", "EXP")[0])   # unitario fraccional
        self.assertFalse(bl.register_shrinkage(pid, "5", "EXP")[0])     # más que el stock (quedan 2)
        self.assertFalse(bl.register_shrinkage(pid, "1", "  ")[0])      # sin motivo
        self.assertEqual(bl.to_decimal_qty(get(Product, pid).stock), D("2.000"))


class CustomerManagementTests(unittest.TestCase):
    def test_save_customer_create_edit_and_validations(self):
        dni = uuid.uuid4().hex[:10]
        ok, msg, c = bl.save_customer({"name": "Marta", "dni_cuit": dni, "phone": "3855000000", "email": "marta@mail.com",
                                       "address": "Belgrano 100", "credit_limit": "50000,00"})
        self.assertTrue(ok, msg)
        self.assertEqual(c["limite_credito"], D("50000.00"))
        self.assertIsNotNone(DCustomer.objects.get(pk=c["id"]).created_at)
        ok, msg, c2 = bl.save_customer({"id": c["id"], "name": "Marta G.", "dni_cuit": dni, "credit_limit": "0"})
        self.assertTrue(ok, msg)
        self.assertEqual((c2["full_name"], c2["limite_credito"]), ("Marta G.", D("0.00")))
        self.assertFalse(bl.save_customer({"name": "Otra", "dni_cuit": dni})[0])           # DNI duplicado
        self.assertFalse(bl.save_customer({"name": "X", "email": "no-es-mail"})[0])
        self.assertFalse(bl.save_customer({"name": "X", "credit_limit": "-1"})[0])
        self.assertFalse(bl.save_customer({"name": "X", "dni_cuit": bl.FINAL_CONSUMER_DNI})[0])  # DNI reservado

    def test_final_consumer_keeps_credit_block(self):
        cf = DCustomer.objects.get(dni_cuit=bl.FINAL_CONSUMER_DNI)
        ok, msg, _ = bl.save_customer({"id": cf.id, "name": cf.full_name, "dni_cuit": bl.FINAL_CONSUMER_DNI, "credit_limit": "1000"})
        self.assertFalse(ok)
        self.assertIn("Consumidor Final", msg)
        ok, msg, _ = bl.save_customer({"id": cf.id, "name": cf.full_name, "dni_cuit": "12345678", "credit_limit": "0"})
        self.assertFalse(ok)
        ok, msg, data = bl.save_customer({"id": cf.id, "name": "Consumidor Final", "dni_cuit": bl.FINAL_CONSUMER_DNI,
                                          "phone": "0", "credit_limit": "0"})
        self.assertTrue(ok, msg)
        self.assertTrue(data["is_final_consumer"])


class DebtPaymentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.user_id = DUser.objects.create_user(f"cobra{uuid.uuid4().hex[:6]}").id

    def test_payment_single_ledger_entry_cash_session_and_receipt(self):
        cust = make_customer(limit="10000")
        pid = make_product("Asado", weighable=True, stock="10", price="8000")
        ok, res = bl.process_sale([{"id": pid, "price": 8000, "quantity": 1}], cust,
                                  payments=[{"payment_method_id": PM["Cuenta Corriente"], "amount": 8000}], user_id=self.user_id)
        self.assertTrue(ok, res)
        # Sin caja abierta no se recibe efectivo
        self.assertFalse(bl.register_debt_payment(cust, "1000", PM["Efectivo"], user_id=self.user_id)[0])
        bl.open_cash_session("0", user_id=self.user_id)
        self.assertFalse(bl.register_debt_payment(cust, "9000", PM["Efectivo"], user_id=self.user_id)[0])        # supera la deuda
        self.assertFalse(bl.register_debt_payment(cust, "0", PM["Efectivo"], user_id=self.user_id)[0])
        self.assertFalse(bl.register_debt_payment(cust, "100", PM["Cuenta Corriente"], user_id=self.user_id)[0])
        ok, msg, cash = bl.register_debt_payment(cust, "3000,50", PM["Efectivo"], notes="a cuenta", user_id=self.user_id)
        self.assertTrue(ok, msg)
        ok, msg, digital = bl.register_debt_payment(cust, "1000", PM["Transferencia / QR"], user_id=self.user_id)
        self.assertTrue(ok, msg)
        self.assertEqual((cash["balance_before"], cash["balance_after"], digital["balance_after"]), (D("8000.00"), D("4999.50"), D("3999.50")))
        session = bl.get_open_cash_session(user_id=self.user_id)
        self.assertEqual((session["total_sales_cash"], session["total_sales_digital"]), (D("3000.50"), D("1000.00")))
        credits = DCustomer.objects.get(pk=cust).account_ledger.filter(entry_type="CREDIT")
        self.assertEqual(credits.count(), 2)  # un asiento por pago, sin duplicados
        self.assertEqual(credits.get(reference__startswith=f"Pago #{cash['payment_id']}").balance_after, D("4999.50"))
        # Recibo
        text = ticket_printer.build_payment_receipt_text(cash, 32)
        self.assertTrue(all(len(line) <= 32 for line in text.splitlines()))
        self.assertIn("SALDO ACTUAL", text)
        self.assertIn("$4999.50", text)
        path = ticket_printer.save_payment_receipt(cash["payment_id"], text, folder=Path(TMP_DIR.name) / "tickets")
        self.assertEqual(path.name, f"recibo_pago_{cash['payment_id']:06d}.txt")
        # Estado de cuenta: venta fiada + 2 pagos
        st = bl.get_customer_account_statement(cust)
        self.assertEqual([e["type"] for e in st["entries"]], ["DEBT", "CREDIT", "CREDIT"])
        self.assertEqual((st["total_debit"], st["total_credit"], st["balance"]), (D("8000.00"), D("4000.50"), D("3999.50")))
        self.assertEqual(bl.register_debt_payment(make_customer(), "1", PM["Efectivo"], user_id=self.user_id)[0], False)  # sin deuda


class SupplierTests(unittest.TestCase):
    def test_save_supplier(self):
        name = f"Frigorífico {uuid.uuid4().hex[:5]}"
        ok, msg, sup = bl.save_supplier({"name": name, "phone": "3854000000", "email": "ventas@frigo.com", "address": "Ruta 9"})
        self.assertTrue(ok, msg)
        self.assertIn(name, [s["name"] for s in bl.get_suppliers_list()])  # nutre el combo de Ingreso de Mercadería
        ok, msg, edited = bl.save_supplier({"id": sup["id"], "name": name + " SA", "email": ""})
        self.assertTrue(ok, msg)
        self.assertEqual((DSupplier.objects.get(pk=sup["id"]).name, edited["email"]), (name + " SA", None))
        self.assertFalse(bl.save_supplier({"name": (name + " sa").upper()})[0])  # duplicado
        self.assertFalse(bl.save_supplier({"name": " "})[0])
        self.assertFalse(bl.save_supplier({"name": "X", "email": "mal"})[0])


class SaleReturnTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.user_id = DUser.objects.create_user(f"dev{uuid.uuid4().hex[:6]}").id
        bl.open_cash_session("10000", user_id=cls.user_id)

    def sell(self, payments, customer=None):
        self.pollo = make_product("Pollo", weighable=True, stock="10", price="4000")
        self.coca = make_product("Coca", stock="10", price="1500")
        ok, res = bl.process_sale([{"id": self.pollo, "price": 4000, "quantity": "1,250"}, {"id": self.coca, "price": 1500, "quantity": 2}],
                                  customer, payments=payments, user_id=self.user_id)  # $8000
        self.assertTrue(ok, res)
        info = bl.get_sale_return_info(res["sale_id"])
        self.lines = {l["product_id"]: l["detail_id"] for l in info["lines"]}
        return res["sale_id"]

    def test_partial_returns_restock_and_cap_cumulative(self):
        sale_id = self.sell([{"payment_method_id": PM["Efectivo"], "amount": 8000}])
        cash_before = bl.get_open_cash_session(user_id=self.user_id)["total_sales_cash"]
        ok, msg, r1 = bl.register_sale_return(sale_id, [{"detail_id": self.lines[self.pollo], "quantity": "0,500"}],
                                              reason="Mal estado", user_id=self.user_id)
        self.assertTrue(ok, msg)
        self.assertEqual((r1["total"], r1["credited"], r1["refund"], r1["refund_method"]), (D("2000.00"), D("0.00"), D("2000.00"), "Efectivo"))
        self.assertEqual(bl.to_decimal_qty(get(Product, self.pollo).stock), D("9.250"))  # 10 - 1.250 + 0.500
        self.assertEqual(bl.get_open_cash_session(user_id=self.user_id)["total_sales_cash"], cash_before - D("2000.00"))
        info = bl.get_sale_return_info(sale_id)
        pollo_line = next(l for l in info["lines"] if l["product_id"] == self.pollo)
        self.assertEqual((pollo_line["returned"], pollo_line["remaining"]), (D("0.500"), D("0.750")))
        # Acumulado: ya se devolvió 0.500 de 1.250 -> no se pueden devolver 0.800
        ok, msg, _ = bl.register_sale_return(sale_id, [{"detail_id": self.lines[self.pollo], "quantity": "0,800"}], user_id=self.user_id)
        self.assertFalse(ok)
        self.assertIn("quedan 0.750 kg", msg)
        self.assertFalse(bl.register_sale_return(sale_id, [{"detail_id": self.lines[self.coca], "quantity": "1,5"}], user_id=self.user_id)[0])
        self.assertFalse(bl.register_sale_return(sale_id, [], user_id=self.user_id)[0])
        self.assertFalse(bl.register_sale_return(999999, [], user_id=self.user_id)[0])
        # Django ve la devolución (stock/movimiento con la misma referencia)
        self.assertTrue(DProduct.objects.get(pk=self.pollo).movements.filter(movement_type="IN", reference=f"Devolución #{r1['return_id']}").exists())
        self.assertEqual(DSaleReturn.objects.get(pk=r1["return_id"]).details.get().quantity, D("0.500"))

    def test_credit_capped_by_fiado_and_rest_refunded(self):
        cust = make_customer(limit="0")
        sale_id = self.sell([{"payment_method_id": PM["Efectivo"], "amount": 6500},
                             {"payment_method_id": PM["Cuenta Corriente"], "amount": 1500}], customer=cust)
        ok, msg, r1 = bl.register_sale_return(sale_id, [{"detail_id": self.lines[self.coca], "quantity": 1}],
                                              refund_method_id=PM["Transferencia / QR"], user_id=self.user_id)
        self.assertTrue(ok, msg)
        self.assertEqual((r1["credited"], r1["refund"]), (D("1500.00"), D("0.00")))
        ok, msg, r2 = bl.register_sale_return(sale_id, [{"detail_id": self.lines[self.pollo], "quantity": "1,250"}],
                                              refund_method_id=PM["Transferencia / QR"], user_id=self.user_id)
        self.assertTrue(ok, msg)
        self.assertEqual((r2["total"], r2["credited"], r2["refund"], r2["refund_method"]),
                         (D("5000.00"), D("0.00"), D("5000.00"), "Transferencia / QR"))  # lo fiado ya se acreditó
        self.assertEqual(bl.to_decimal_money(get(Customer, cust).balance), D("0.00"))
        self.assertEqual(bl.capped_return_credit("5000", "1500", "1500"), D("0.00"))
        sales = {s["id"]: s for s in bl.get_recent_sales()}
        self.assertEqual(sales[sale_id]["returned"], D("6500.00"))

    def test_cash_refund_requires_open_cash_session(self):
        uid = DUser.objects.create_user(f"sincaja{uuid.uuid4().hex[:6]}").id
        pid = make_product("Pan", price="100")
        ok, res = bl.process_sale([{"id": pid, "price": 100, "quantity": 1}], None, PM["Efectivo"], user_id=uid)
        detail_id = bl.get_sale_return_info(res["sale_id"])["lines"][0]["detail_id"]
        ok, msg, _ = bl.register_sale_return(res["sale_id"], [{"detail_id": detail_id, "quantity": 1}], user_id=uid)
        self.assertFalse(ok)
        self.assertIn("Abra la caja", msg)
        ok, msg, _ = bl.register_sale_return(res["sale_id"], [{"detail_id": detail_id, "quantity": 1}],
                                             refund_method_id=PM["Transferencia / QR"], user_id=uid)
        self.assertTrue(ok, msg)


class CashSessionSyncTests(unittest.TestCase):
    """SPEC-FIN-07: escritorio y web calculan el mismo arqueo; la web ya no pierde los reintegros."""

    def test_formula(self):
        from core.cash import cash_session_totals
        t = cash_session_totals(initial_amount="1000",
                                sale_payments=[("Efectivo", False, "2000"), ("Transferencia / QR", True, "700"), ("Cuenta Corriente", False, "800")],
                                debt_payments=[("Efectivo", False, "300"), ("Tarjeta Débito", True, "500")],
                                refunds=[("Efectivo", False, "250")], expenses=["100"])
        self.assertEqual((t["total_sales_cash"], t["total_sales_digital"], t["total_expenses"], t["expected_final_amount"]),
                         (D("2050.00"), D("1200.00"), D("100.00"), D("2950.00")))

    def test_desktop_operations_then_web_dashboard_keeps_refund(self):
        from django.test import Client
        user = DUser.objects.create_user(f"cruce{uuid.uuid4().hex[:6]}", password="x", is_staff=True)
        uid = user.id
        self.assertTrue(bl.open_cash_session("1000", user_id=uid)[0])
        pid = make_product("Pollo", weighable=True, stock="20", price="4000")
        cust = make_customer(limit="0")
        ok, sale = bl.process_sale([{"id": pid, "price": 4000, "quantity": "0,500"}], None, PM["Efectivo"], user_id=uid)   # +2000 efectivo
        self.assertTrue(ok, sale)
        ok, fiado = bl.process_sale([{"id": pid, "price": 4000, "quantity": "0,200"}], cust, PM["Cuenta Corriente"], user_id=uid)  # 800 fiado
        self.assertTrue(ok, fiado)
        self.assertTrue(bl.register_debt_payment(cust, "500", PM["Transferencia / QR"], user_id=uid)[0])                 # +500 digital
        detail = bl.get_sale_return_info(sale["sale_id"])["lines"][0]["detail_id"]
        ok, msg, ret = bl.register_sale_return(sale["sale_id"], [{"detail_id": detail, "quantity": "0,100"}], user_id=uid)  # -400 efectivo
        self.assertTrue(ok, msg)
        self.assertTrue(bl.add_cash_expense("100", "Hielo", user_id=uid)[0])
        self.assertEqual(DSaleReturn.objects.get(pk=ret["return_id"]).refund_method.name, "Efectivo")

        desktop = bl.get_open_cash_session(user_id=uid)
        self.assertEqual((desktop["total_sales_cash"], desktop["total_sales_digital"], desktop["total_expenses"]),
                         (D("1600.00"), D("500.00"), D("100.00")))

        # Abrir el dashboard web (antes recalculaba y PISABA la caja sin el reintegro de 400)
        client = Client()
        client.force_login(user)
        from django.urls import reverse
        response = client.get(reverse("finance:cash_dashboard"))
        self.assertEqual(response.status_code, 200)
        from finance.models import CashSession as DCashSession
        web = DCashSession.objects.get(user=user, is_open=True)
        self.assertEqual((web.total_sales_cash, web.total_sales_digital, web.total_expenses, web.expected_final_amount),
                         (D("1600.00"), D("500.00"), D("100.00"), D("2500.00")))  # 1000 + 2000 - 400 - 100

        # Y el escritorio, al volver, ve lo mismo y cierra sin diferencia
        self.assertEqual(bl.get_open_cash_session(user_id=uid)["total_sales_cash"], D("1600.00"))
        ok, msg = bl.close_cash_session("2500", user_id=uid)
        self.assertTrue(ok, msg)
        self.assertIn("Diferencia: $0.00", msg)


class BackupTests(unittest.TestCase):
    """SPEC-BACKUP-01: respaldo diario idempotente, rotación de 15 y copia manual / externa."""

    def test_daily_backup_created_at_startup_and_idempotent(self):
        import sqlite3
        from core import backup_manager
        folder = backup_manager.backups_dir(TMP_DB)
        daily = [p for p in backup_manager.list_backups(folder)]
        self.assertGreaterEqual(len(daily), 1)  # lo creó prepare_database() al iniciar
        self.assertRegex(daily[0].name, r"^impulso_smart_\d{8}_\d{6}(_\d+)?\.sqlite3$")
        self.assertIsNone(backup_manager.ensure_daily_backup(TMP_DB))  # ya hay uno de hoy
        con = sqlite3.connect(daily[0])
        self.assertEqual(con.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertGreater(con.execute("select count(*) from finance_paymentmethod").fetchone()[0], 0)
        con.close()
        self.assertIsNone(backup_manager.ensure_daily_backup(":memory:"))

    def test_rotation_keeps_newest_15(self):
        from core import backup_manager
        folder = Path(TMP_DIR.name) / "rotacion"
        folder.mkdir()
        for day in range(1, 19):
            (folder / f"impulso_smart_202601{day:02d}_120000.sqlite3").write_bytes(b"x")
        (folder / "pre_migracion_20250101_000000.sqlite3").write_bytes(b"x")  # no entra en la rotación
        (folder / "notas.txt").write_text("no tocar")
        removed = backup_manager.rotate_backups(folder, keep=15)
        self.assertEqual(sorted(p.name for p in removed),
                         [f"impulso_smart_202601{d:02d}_120000.sqlite3" for d in (1, 2, 3)])
        self.assertEqual(len(backup_manager.list_backups(folder)), 15)
        self.assertTrue((folder / "pre_migracion_20250101_000000.sqlite3").exists() and (folder / "notas.txt").exists())
        # Una copia nueva desplaza a la más vieja
        path = backup_manager.ensure_daily_backup(TMP_DB, folder)
        self.assertIsNotNone(path)
        self.assertEqual(len(backup_manager.list_backups(folder)), 15)
        self.assertFalse((folder / "impulso_smart_20260104_120000.sqlite3").exists())

    def test_manual_and_external_backup(self):
        before = len(bl.get_backup_status()["backups"])
        ok, msg, data = bl.create_manual_backup()
        self.assertTrue(ok, msg)
        self.assertEqual(len(bl.get_backup_status()["backups"]), before + 1)
        usb = Path(TMP_DIR.name) / "pendrive"
        ok, msg, data = bl.create_manual_backup(usb)
        self.assertTrue(ok, msg)
        self.assertEqual(data["path"].parent, usb)
        blocked = Path(TMP_DIR.name) / "es_un_archivo"
        blocked.write_text("x")  # destino inválido: un archivo, no una carpeta
        ok, msg, _ = bl.create_manual_backup(blocked)
        self.assertFalse(ok)
        self.assertIn("No se pudo crear la copia de seguridad", msg)

if __name__ == "__main__":
    result = unittest.main(verbosity=2, exit=False).result
    # Cerrar conexiones antes de borrar la base temporal (en Windows no se puede borrar un archivo abierto)
    from django.db import connections
    connections.close_all()
    db_manager.db.engine.dispose()
    TMP_DIR.cleanup()
    sys.exit(0 if result.wasSuccessful() else 1)
