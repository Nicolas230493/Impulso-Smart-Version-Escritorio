"""Exportación del resumen de rentabilidad y auditoría de caja a CSV o Excel."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SUMMARY_ROWS = [
    ("Ventas brutas", "gross_sales"),
    ("Devoluciones", "returns"),
    ("Ventas netas", "total_sales"),
    ("CMV (costo de mercadería vendida)", "cmv"),
    ("Margen bruto", "gross_profit"),
    ("Margen bruto %", "gross_margin_pct"),
    ("Gastos de caja", "expenses"),
    ("Mermas / bajas valorizadas", "shrinkage_cost"),
    ("GANANCIA NETA REAL", "net_profit"),
    ("Cantidad de ventas", "sales_count"),
    ("Ajustes por desposte dinámico", "auto_adjustments_count"),
]


def _fmt_date(d):
    return d.strftime("%d/%m/%Y %H:%M") if d else ""


def _sections(summary):
    """Secciones del reporte como (título, encabezados, filas) para ambos formatos."""
    period = f"{_fmt_date(summary['start'])} a {_fmt_date(summary['end'])}"
    head = [("Período", period)]
    if summary.get("cash_session_id"):
        head.append(("Turno de caja", f"#{summary['cash_session_id']}"))
    yield "Resumen", ("Concepto", "Valor"), head + [(label, summary[key]) for label, key in SUMMARY_ROWS]
    yield "Pagos", ("Medio de pago", "Monto"), list(summary["payments_breakdown"].items())
    yield "Despostes", ("Fecha", "Producto", "Cantidad ajustada", "Referencia"), [
        (_fmt_date(a["date"]), a["product"], a["quantity"], a["reference"]) for a in summary["auto_adjustments"]]
    yield "Mermas", ("Fecha", "Producto", "Cantidad", "Costo unit.", "Costo total", "Motivo"), [
        (_fmt_date(l["date"]), l["product"], l["quantity"], l["unit_cost"], l["cost"], l["reason"]) for l in summary["losses"]]


def export_summary(summary, path):
    """Guarda el resumen en .xlsx (una hoja por sección) o .csv (secciones separadas). Devuelve la ruta."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for title, headers, rows in _sections(summary):
            ws = wb.create_sheet(title)
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in rows:
                ws.append([float(v) if hasattr(v, "quantize") else v for v in row])
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(12, *(len(str(c.value or "")) + 2 for c in col))
        wb.save(path)
    else:
        # ";" y BOM para que Excel en español lo abra en columnas directamente
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            for title, headers, rows in _sections(summary):
                writer.writerow([title.upper()])
                writer.writerow(headers)
                writer.writerows(rows)
                writer.writerow([])
    return path
